from io import BytesIO

from docx import Document as DocxDocument

from app.core.exceptions import DocumentTextExtractionError
from app.services.document_text_extractor import DocumentTextExtractor
import pytest


def test_plain_text_uses_markdown_headings_as_sections() -> None:
    content = """# Politica de pagos

El pago se realiza en cinco dias.

# Excepciones

No aplica a proveedores internacionales.
""".encode("utf-8")

    sections = DocumentTextExtractor().extract("politica.txt", content)

    assert [section.title for section in sections] == ["Politica de pagos", "Excepciones"]
    assert "cinco dias" in sections[0].text
    assert sections[0].page_number is None
    assert "proveedores internacionales" in sections[1].text


def test_plain_text_without_headings_is_one_section() -> None:
    sections = DocumentTextExtractor().extract("nota.txt", "Texto operativo sin títulos.".encode("utf-8"))

    assert len(sections) == 1
    assert sections[0].title is None
    assert sections[0].text == "Texto operativo sin títulos."


def test_docx_uses_heading_styles_and_keeps_table_text() -> None:
    document = DocxDocument()
    document.add_heading("Politica de pagos", level=1)
    document.add_paragraph("El pago se realiza en cinco dias.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Canal"
    table.rows[0].cells[1].text = "Transferencia"
    buffer = BytesIO()
    document.save(buffer)

    sections = DocumentTextExtractor().extract("politica.docx", buffer.getvalue())

    assert sections[0].title == "Politica de pagos"
    assert "cinco dias" in sections[0].text
    assert "Transferencia" in sections[0].text


def test_pdf_keeps_page_number_and_readable_text() -> None:
    sections = DocumentTextExtractor().extract("politica.pdf", _build_pdf(["1. Politica de pagos", "El pago se realiza en cinco dias."]))

    assert sections
    combined = " ".join(section.text for section in sections)
    assert "Politica de pagos" in combined or sections[0].title == "1 Politica de pagos"
    assert sections[0].page_number == 1


def test_empty_text_file_fails() -> None:
    with pytest.raises(DocumentTextExtractionError):
        DocumentTextExtractor().extract("vacio.txt", b"   \n")


def _build_pdf(lines: list[str]) -> bytes:
    content_lines = []
    y_position = 740
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content_lines.append(f"BT /F1 14 Tf 72 {y_position} Td ({escaped}) Tj ET")
        y_position -= 24
    stream = "\n".join(content_lines).encode("latin-1")
    objects = [
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
        b"4 0 obj\n<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream\nendobj\n",
        b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n",
    ]
    header = b"%PDF-1.4\n"
    offset = len(header)
    offsets = []
    body = b""
    for pdf_object in objects:
        offsets.append(offset)
        body += pdf_object
        offset += len(pdf_object)
    xref = b"xref\n" + f"0 {len(objects) + 1}\n".encode("ascii") + b"0000000000 65535 f \n"
    for object_offset in offsets:
        xref += f"{object_offset:010d} 00000 n \n".encode("ascii")
    trailer = (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode("ascii")
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(offset).encode("ascii")
        + b"\n%%EOF\n"
    )
    return header + body + xref + trailer
